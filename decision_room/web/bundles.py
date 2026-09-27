"""Resumable, business-scoped folder uploads and Excel worksheet preparation."""
import fcntl
import json
import os
import shutil
from contextlib import contextmanager
from datetime import date, datetime, time
from pathlib import Path, PurePosixPath
from tempfile import TemporaryDirectory
from uuid import UUID
from zipfile import ZipFile

from openpyxl import load_workbook

from ..database import connect
from ..storage import Storage, digest
from . import dossier
from .errors import WebError, identifier

CHUNK_BYTES = 8 * 1024 * 1024
MANIFEST_BYTES = 2 * 1024 * 1024


def relative_name(value):
    if not isinstance(value, str) or not value or len(value) > 1024 or '\\' in value:
        raise WebError('Revisa los nombres de los archivos de la carpeta.')
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in ('', '.', '..') for part in value.split('/')) or any(ord(c) < 32 for c in value):
        raise WebError('Revisa los nombres de los archivos de la carpeta.')
    if path.suffix.lower() not in ('.csv', '.xlsx'):
        raise WebError('La carpeta solo puede incluir CSV y Excel .xlsx.')
    return value


def folder(ws, upload_id):
    business = ws.business_id()
    return Storage(ws.config.storage).path(business, f'{business}/bundle-uploads/{identifier(upload_id)}')


@contextmanager
def locked(directory):
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    with (directory / '.lock').open('a+b') as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(lock, fcntl.LOCK_UN)


def _read(directory):
    try:
        return json.loads((directory / 'manifest.json').read_text())
    except (FileNotFoundError, ValueError):
        raise WebError('No encontramos esta entrega. Vuelve a seleccionar la carpeta.', 404) from None


def _state(directory, manifest):
    return {'id': manifest['request_key'], 'files': [
        {'path': item['path'], 'size': item['size'],
         'uploaded': (directory / f'{index}.part').stat().st_size if (directory / f'{index}.part').exists() else 0}
        for index, item in enumerate(manifest['files'])],
        'result': json.loads((directory / 'result.json').read_text()) if (directory / 'result.json').exists() else None}


def start(ws, data):
    business = dossier.guard(ws, data)
    allowed = {'business_id', 'request_key', 'title', 'mode', 'previous_id',
               'period_from', 'period_until', 'files'}
    if set(data) - allowed:
        raise WebError('La descripción de la entrega contiene campos desconocidos.')
    key = str(identifier(data.get('request_key')))
    files = data.get('files')
    if not isinstance(files, list) or not files or len(files) > ws.config.max_files:
        raise WebError('Selecciona una carpeta con archivos CSV o Excel.')
    listed = []
    for item in files:
        if not isinstance(item, dict) or set(item) != {'path', 'size'}:
            raise WebError('La lista de archivos no es válida.')
        path = relative_name(item['path'])
        size = item['size']
        if type(size) is not int or not 0 < size <= ws.config.max_file_bytes:
            raise WebError('Un archivo supera el límite de 2 GB.')
        listed.append({'path': path, 'size': size})
    if len({item['path'].casefold() for item in listed}) != len(listed):
        raise WebError('Hay archivos con la misma ruta en la carpeta.')
    if sum(item['size'] for item in listed) > ws.config.max_batch_bytes:
        raise WebError('La carpeta supera el límite total de 2 GB.', 413)
    metadata = {k: v for k, v in data.items() if k != 'files'}
    if len(json.dumps(metadata, ensure_ascii=False)) > 24_000:
        raise WebError('La descripción de la entrega es demasiado larga.')
    manifest = {**metadata, 'business_id': str(business), 'request_key': key,
                'files': sorted(listed, key=lambda x: x['path'])}
    directory = folder(ws, key)
    with locked(directory):
        path = directory / 'manifest.json'
        if path.exists():
            if _read(directory) != manifest:
                raise WebError('Esta entrega ya existe con otros archivos o datos.', 409)
        else:
            with path.open('x') as stream:
                json.dump(manifest, stream, ensure_ascii=False)
                stream.flush()
                os.fsync(stream.fileno())
        return _state(directory, manifest)


def status(ws, upload_id):
    directory = folder(ws, upload_id)
    with locked(directory):
        return _state(directory, _read(directory))


def chunk(ws, upload_id, index, offset, data):
    directory = folder(ws, upload_id)
    with locked(directory):
        manifest = _read(directory)
        if (directory / 'result.json').exists():
            raise WebError('La entrega ya está terminada.', 409)
        if not isinstance(index, int) or not 0 <= index < len(manifest['files']):
            raise WebError('Archivo no encontrado.', 404)
        size = manifest['files'][index]['size']
        if not data or len(data) > CHUNK_BYTES or type(offset) is not int or offset < 0 or offset + len(data) > size:
            raise WebError('Este fragmento supera el tamaño del archivo.', 413)
        target = directory / f'{index}.part'
        current = target.stat().st_size if target.exists() else 0
        if offset > current:
            raise WebError('Falta un fragmento anterior. Reanuda desde el avance indicado.', 409)
        if offset < current:
            if offset + len(data) > current:
                raise WebError('Los fragmentos se solapan. Consulta el avance y reintenta.', 409)
            with target.open('rb') as stream:
                stream.seek(offset)
                if stream.read(len(data)) != data:
                    raise WebError('El archivo seleccionado ha cambiado. Inicia otra entrega.', 409)
        else:
            with target.open('ab') as stream:
                stream.write(data)
                stream.flush()
                os.fsync(stream.fileno())
        return {'uploaded': target.stat().st_size, 'size': size}


def _cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _csv_record(values):
    # Preserve blank Excel cells as CSV NULL and explicit empty strings as "".
    return ','.join('' if value is None else '"' + _cell(value).replace('"', '""') + '"'
                    for value in values) + '\n'


def _excel(workbook_path, relative, converted, config, budget):
    with ZipFile(workbook_path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 4_000_000_000:
            raise ValueError('El Excel supera el límite de descompresión de 4 GB.')
    book = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    outputs = []
    try:
        for number, sheet in enumerate(book.worksheets):
            if number >= 256:
                raise ValueError('El Excel contiene demasiadas hojas.')
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None or not any(value is not None for value in header):
                continue
            if len(header) > config.max_columns or any(value is None or not str(value).strip() for value in header):
                raise ValueError(f'La hoja {sheet.title} necesita encabezados válidos.')
            name = f'{workbook_path.stem}-{number:03d}.csv'
            path = converted / name
            with path.open('w', encoding='utf-8', newline='') as stream:
                stream.write(_csv_record([str(value) for value in header]))
                for count, row in enumerate(rows, 1):
                    if count > config.max_rows:
                        raise ValueError(f'La hoja {sheet.title} supera el límite de filas.')
                    if not any(value is not None for value in row):
                        continue
                    stream.write(_csv_record(row))
                    if count % 1000 == 0 and stream.tell() + budget[0] > config.max_batch_bytes:
                        raise ValueError('Los datos preparados superan 2 GB. Reduce la entrega.')
            budget[0] += path.stat().st_size
            if budget[0] > config.max_batch_bytes:
                raise ValueError('Los datos preparados superan 2 GB. Reduce la entrega.')
            outputs.append((path, f'{relative} [{sheet.title}].csv'))
    finally:
        book.close()
    return outputs


def finish(ws, upload_id):
    directory = folder(ws, upload_id)
    with locked(directory):
        manifest = _read(directory)
        if (directory / 'result.json').exists():
            return json.loads((directory / 'result.json').read_text())
        paths = []
        for index, item in enumerate(manifest['files']):
            path = directory / f'{index}.part'
            if not path.exists() or path.stat().st_size != item['size']:
                raise WebError('La carpeta todavía no se ha subido por completo.', 409)
            paths.append(path)
        with TemporaryDirectory(prefix='dr-bundle-') as temp:
            converted = Path(temp)
            inputs, names, budget = [], {}, [sum(p.stat().st_size for p, f in zip(paths, manifest['files']) if f['path'].lower().endswith('.csv'))]
            originals = []
            storage = Storage(ws.config.storage)
            for index, (path, item) in enumerate(zip(paths, manifest['files'])):
                original_path = converted / f'original-{index}{PurePosixPath(item["path"]).suffix.lower()}'
                shutil.copyfile(path, original_path)
                saved = storage.capture(ws.business_id(), original_path, ws.config.max_file_bytes,
                                        original_name=item['path'])
                originals.append((item['path'], saved))
                if original_path.suffix == '.csv':
                    inputs.append(original_path)
                    names[original_path.resolve()] = item['path']
                else:
                    try:
                        sheets = _excel(original_path, item['path'], converted, ws.config, budget)
                    except Exception:
                        # Give this original its own failed source so other files
                        # in the same delivery can still become usable tables.
                        failed = converted / f'unreadable-{index}.csv'
                        failed.write_bytes(b'\xff' + saved['sha256'].encode())
                        label = f'{item["path"]} [Excel no legible].csv'
                        inputs.append(failed)
                        names[failed.resolve()] = label
                        continue
                    for sheet_path, label in sheets:
                        inputs.append(sheet_path)
                        names[sheet_path.resolve()] = label
            if not inputs:
                raise WebError('No hay hojas con datos en la carpeta.')
            result = dossier.upload_paths(ws, manifest, inputs, names=names,
                                          original_signature=[(name, item['sha256']) for name, item in originals])
            with connect(ws.config) as db, db.transaction():
                for name, item in originals:
                    db.execute('''INSERT INTO dataset_bundle_files(business_id,analysis_id,relative_path,sha256,byte_count,original_key)
                        VALUES (%s,%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING''',
                        (ws.business_id(), UUID(result['analysis_id']), name, item['sha256'],
                         item['byte_count'], item['original_key']))
            with (directory / 'result.json').open('w') as stream:
                json.dump(result, stream, default=str)
                stream.flush()
                os.fsync(stream.fileno())
            for path in paths:
                path.unlink(missing_ok=True)
            return result


def original_location(ws, analysis_id, relative_path):
    with connect(ws.config) as db:
        row = db.execute('''SELECT original_key,sha256 FROM dataset_bundle_files
            WHERE business_id=%s AND analysis_id=%s AND relative_path=%s''',
            (ws.business_id(), identifier(analysis_id), relative_name(relative_path))).fetchone()
    if not row:
        raise WebError('Archivo no encontrado.', 404)
    path = Storage(ws.config.storage).path(ws.business_id(), row['original_key'])
    if digest(path) != row['sha256']:
        raise WebError('El archivo original ha cambiado.', 409)
    return PurePosixPath(relative_path).name, path
