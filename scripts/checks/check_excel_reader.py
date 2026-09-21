"""Independent openpyxl read checks of each exported data sheet."""
from datetime import datetime
from decimal import Decimal
import json
import openpyxl
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts/datasets'))

from convert_bacpac import OUT
from prepare_excel_plan import typed_rows

def main():
    plan=json.loads((OUT/"metadata/excel-plan.json").read_text())
    manifest=json.loads((OUT/"manifest.json").read_text())
    checks=[]
    for group,sheets in plan["groups"].items():
        file=OUT/"excel"/f"WideWorldImporters-{group}.xlsx"
        wb=openpyxl.load_workbook(file,read_only=True,data_only=False)
        for spec in sheets:
            ws=wb[spec["sheet"]]
            assert ws.max_row==spec["rows"]+1
            assert ws.max_column==len(spec["columns"])
            rows=ws.iter_rows(min_row=1,max_row=min(9,spec["rows"]+1),values_only=True)
            assert list(next(rows))==[c["name"]for c in spec["columns"]]
            for row,source in zip(rows,spec["sample"]):
                for actual,expected,c in zip(row,source,spec["columns"]):
                    if expected is None:
                        assert actual is None,(spec["sheet"],c["name"],actual)
                    elif len(expected)>32767:
                        assert actual.startswith("../large-values/")
                    elif c["type"]=="date":
                        assert actual.date().isoformat()==expected
                    elif c["type"]=="bit":
                        assert isinstance(actual,bool) and actual==(expected=="1")
                    elif c["type"] in ("int","bigint","decimal") and not c["name"].endswith("ID") and len(Decimal(expected).as_tuple().digits)<=15:
                        assert Decimal(str(actual))==Decimal(expected)
                    else:
                        # openpyxl reads an explicit empty inline string as ''.
                        assert actual==expected,(spec["sheet"],c["name"],actual,expected)
            checks.append({"workbook":file.name,"sheet":spec["sheet"],"rows":spec["rows"],"sample_rows_checked":len(spec["sample"])})
        wb.close()
        print("Independent Excel reader passed:",file.name,flush=True)
    (OUT/"metadata/independent-reader-checks.json").write_text(json.dumps(checks,indent=2)+"\n")

if __name__=="__main__":
    main()
