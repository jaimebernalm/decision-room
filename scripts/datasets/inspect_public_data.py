from pathlib import Path
from itertools import islice
import json
import zipfile
import openpyxl
import csv
import hashlib
import io

repo = Path(__file__).resolve().parents[2]
root = repo / 'docs/research/2026-09-17'
data = repo / 'data/exploratory/2026-09-17'
root.mkdir(parents=True, exist_ok=True)
with zipfile.ZipFile(data / 'online-retail-ii.zip') as archive:
    names = [n for n in archive.namelist() if n.endswith('.xlsx')]
    assert len(names) == 1, names
    workbook_path = data / Path(names[0]).name
    workbook_path.write_bytes(archive.read(names[0]))

wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
result = {
    'source': 'https://archive.ics.uci.edu/dataset/502/online+retail+ii',
    'source_zip_url': 'https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip',
    'inspection': 'Workbook dimensions, actual headers and first 500 data rows of each sheet. Not a full data-quality audit.',
    'sheets': [],
}
for ws in wb:
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    sample = list(islice(rows, 500))
    result['sheets'].append({
        'sheet': ws.title,
        'dimension_rows_including_header': ws.max_row,
        'dimension_columns': ws.max_column,
        'headers': list(headers),
        'sample_rows_read': len(sample),
        'sample_null_counts': {str(col): sum(row[i] is None for row in sample) for i, col in enumerate(headers)},
    })
wb.close()
(root / 'public_data_inspection.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False, indent=2))

# Read the two smaller CSV sources completely for structural checks only.
# These checks do not validate business conclusions or prediction performance.
sport_path = data / 'sport-services.csv'
sport_text = sport_path.read_text(encoding='utf-8-sig')
sport_dialect = csv.Sniffer().sniff(sport_text[:12000], delimiters=',;\t')
sport_rows = list(csv.reader(io.StringIO(sport_text), sport_dialect))
sport_headers = sport_rows[0]
sport_result = {
    'source': 'https://data.mendeley.com/datasets/yprk4jdgnv/1',
    'filename': sport_path.name,
    'sha256': hashlib.sha256(sport_path.read_bytes()).hexdigest(),
    'delimiter': sport_dialect.delimiter,
    'rows_excluding_header': len(sport_rows) - 1,
    'columns': len(sport_headers),
    'headers': sport_headers,
    'row_width_counts': {str(n): sum(len(row) == n for row in sport_rows[1:]) for n in set(map(len, sport_rows[1:]))},
    'empty_counts': {col: sum(row[i] == '' for row in sport_rows[1:]) for i, col in enumerate(sport_headers)},
}
(root / 'sport_data_inspection.json').write_text(json.dumps(sport_result, ensure_ascii=False, indent=2))

hotel_path = data / 'hotel-original-2.zip'
hotel_result = {
    'source': 'https://doi.org/10.1016/j.dib.2018.11.126',
    'download_url': 'https://ars.els-cdn.com/content/image/1-s2.0-S2352340918315191-mmc2.zip',
    'sha256': hashlib.sha256(hotel_path.read_bytes()).hexdigest(),
    'files': [],
}
with zipfile.ZipFile(hotel_path) as archive:
    for name in archive.namelist():
        if name.startswith('__MACOSX/') or not name.lower().endswith('.csv'):
            continue
        rows = list(csv.reader(io.StringIO(archive.read(name).decode('utf-8-sig'))))
        headers = rows[0]
        hotel_result['files'].append({
            'name': name,
            'rows_excluding_header': len(rows) - 1,
            'columns': len(headers),
            'headers': headers,
            'row_width_counts': {str(n): sum(len(row) == n for row in rows[1:]) for n in set(map(len, rows[1:]))},
        })
(root / 'hotel_data_inspection.json').write_text(json.dumps(hotel_result, ensure_ascii=False, indent=2))
print(json.dumps({'sports_rows': len(sport_rows) - 1, 'sports_columns': len(sport_headers), 'hotel_files': hotel_result['files']}, ensure_ascii=False, indent=2))
